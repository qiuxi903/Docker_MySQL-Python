"""导入模块：CSV / Excel(.xlsx) / SQL 脚本 三种格式导入到学生自己的库。"""
import base64
import csv
import datetime
import io
import re
import time
import zlib

import pymysql

from .config import C
from .db import q_ident, root_conn, user_conn
from .sqlrun import run_sql


class ImportError_(Exception):
    pass


# ---------- 源文件存档（表 ← 导入时的原始文件，可预览/下载） ----------
# 存在内部 meta 库，学生自己的库里不出现任何系统表

def _ensure_sources(cur):
    cur.execute("""
        CREATE TABLE IF NOT EXISTS sources (
            db_name VARCHAR(64) NOT NULL,
            tab VARCHAR(64) NOT NULL,
            filename VARCHAR(255) NOT NULL DEFAULT '',
            kind VARCHAR(16) NOT NULL DEFAULT '',
            content LONGTEXT,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (db_name, tab)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)


def store_source(db, mysql_user, mysql_pass, table, filename, kind, data: bytes):
    if len(data) > 4 * 1024 * 1024:
        data = data[: 4 * 1024 * 1024]
    payload = base64.b64encode(zlib.compress(data)).decode("ascii")
    conn = root_conn(C.META_DB)
    try:
        with conn.cursor() as cur:
            _ensure_sources(cur)
            cur.execute(
                "REPLACE INTO sources (db_name, tab, filename, kind, content) VALUES (%s,%s,%s,%s,%s)",
                (db, table, filename[:250], kind, payload))
    finally:
        conn.close()


def get_source(db, mysql_user, mysql_pass, table):
    """返回 {filename, kind, data(bytes), createdAt} 或 None。"""
    conn = root_conn(C.META_DB)
    try:
        with conn.cursor() as cur:
            try:
                cur.execute("SELECT filename, kind, content, created_at FROM sources "
                            "WHERE db_name=%s AND tab=%s", (db, table))
                row = cur.fetchone()
            except pymysql.MySQLError:
                return None
    finally:
        conn.close()
    if not row:
        return None
    return {
        "filename": row[0],
        "kind": row[1],
        "data": zlib.decompress(base64.b64decode(row[2])),
        "createdAt": str(row[3]),
    }


def drop_source(db, table):
    conn = root_conn(C.META_DB)
    try:
        with conn.cursor() as cur:
            try:
                cur.execute("DELETE FROM sources WHERE db_name=%s AND tab=%s", (db, table))
            except pymysql.MySQLError:
                pass
    finally:
        conn.close()


def _decode(data: bytes) -> str:
    """中文场景常见编码兜底：utf-8(sig) → gbk → gb18030。"""
    for enc in ("utf-8-sig", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    try:
        return data.decode("gb18030")
    except UnicodeDecodeError:
        raise ImportError_("文件编码无法识别，请用 UTF-8 或 GBK 编码保存")


def valid_table_name(name: str) -> bool:
    from .db import valid_ident

    return valid_ident(name)


def _clean_columns(headers: list) -> list[str]:
    cols: list[str] = []
    for i, h in enumerate(headers, 1):
        h = str(h).strip() if h is not None else ""
        h = re.sub(r"[^\w\u4e00-\u9fff-]", "_", h, flags=re.UNICODE)[:60]
        if not h or h.startswith("_") and h.strip("_") == "":
            h = f"col{i}"
        if h[0].isdigit():
            h = "c" + h
        base = h
        k = 2
        while h in cols:
            h = f"{base}_{k}"
            k += 1
        cols.append(h)
    return cols


_INT_RE = re.compile(r"^[+-]?\d{1,18}$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_TIME_RE = re.compile(r"^\d{2}:\d{2}(:\d{2})?$")


def _infer_type(values: list[str]) -> str:
    """根据一列的非空值推断 SQL 类型。"""
    seen = [v for v in values if v is not None and v != ""]
    if not seen:
        return "VARCHAR(255)"
    if all(_INT_RE.match(v) for v in seen):
        return "BIGINT"
    try:
        for v in seen:
            f = float(v)
            if f != f or f in (float("inf"), float("-inf")):
                raise ValueError
        return "DOUBLE"
    except ValueError:
        pass
    if all(_DATE_RE.match(v) for v in seen):
        return "DATE"
    if all(_TIME_RE.match(v) for v in seen):
        return "TIME"
    if max(len(v) for v in seen) > 255:
        return "TEXT"
    return "VARCHAR(255)"


def _rows_from_grid(grid: list[list], has_header: bool):
    """二维数组 → (列名, 类型列表, 行数据)。"""
    if not grid:
        raise ImportError_("文件内容为空")
    if len(grid[0]) > C.IMPORT_COLS_MAX:
        raise ImportError_(f"列数超过上限（{C.IMPORT_COLS_MAX} 列）")
    header = grid[0] if has_header else [f"col{i}" for i in range(1, len(grid[0]) + 1)]
    body = grid[1:] if has_header else grid
    width = max(len(r) for r in grid)
    cols = _clean_columns(list(header) + [""] * (width - len(header)))

    norm: list[list] = []
    for r in body:
        r = list(r) + [None] * (width - len(r))
        out = []
        for v in r:
            if v is None or v == "":
                out.append(None)
            elif isinstance(v, datetime.datetime):
                out.append(v.isoformat(sep=" ")[:19])
            elif isinstance(v, datetime.date):
                out.append(v.isoformat())
            elif isinstance(v, datetime.time):
                out.append(v.isoformat())
            else:
                s = str(v)
                if len(s) > 4000:
                    raise ImportError_("单个单元格内容超过 4000 字符")
                out.append(s)
        norm.append(out)
    if len(norm) > C.IMPORT_ROWS_MAX:
        raise ImportError_(f"行数超过上限（{C.IMPORT_ROWS_MAX} 行）")

    types = []
    for j in range(width):
        types.append(_infer_type([r[j] for r in norm]))
    return cols, types, norm


def _create_and_fill(db, mysql_user, mysql_pass, table, cols, types, rows, overwrite):
    conn = user_conn(db, mysql_user, mysql_pass)
    try:
        with conn.cursor() as cur:
            cur.execute("SHOW TABLES")
            existing = {r[0] for r in cur.fetchall()}
            if table in existing and not overwrite:
                raise ImportError_(f"表 `{table}` 已存在。勾选“覆盖同名表”可以替换它。")
            if table in existing:
                cur.execute(f"DROP TABLE {q_ident(table)}")
            col_defs = ", ".join(f"{q_ident(c)} {t}" for c, t in zip(cols, types))
            cur.execute(
                f"CREATE TABLE {q_ident(table)} ({col_defs}) "
                "ENGINE=InnoDB DEFAULT CHARSET=utf8mb4"
            )
            if rows:
                placeholders = ", ".join(["%s"] * len(cols))
                sql = f"INSERT INTO {q_ident(table)} VALUES ({placeholders})"
                chunk = 500
                for i in range(0, len(rows), chunk):
                    cur.executemany(sql, rows[i : i + chunk])
            conn.commit()
        return len(rows)
    except pymysql.MySQLError as e:
        try:
            conn.rollback()
        except Exception:
            pass
        msg = e.args[1] if len(e.args) > 1 else str(e)
        raise ImportError_(f"导入失败：{msg}") from e
    finally:
        conn.close()


def import_csv(db, mysql_user, mysql_pass, table, data: bytes, has_header: bool, overwrite: bool,
               source_name: str = "") -> dict:
    text = _decode(data)
    # 取首行探测分隔符
    first_line = text.splitlines()[0] if text.splitlines() else ""
    delim = max([",", ";", "\t", "|"], key=first_line.count)
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    grid = [row for row in reader if any(cell.strip() for cell in row)]
    t0 = time.monotonic()
    cols, types, rows = _rows_from_grid(grid, has_header)
    n = _create_and_fill(db, mysql_user, mysql_pass, table, cols, types, rows, overwrite)
    try:
        store_source(db, mysql_user, mysql_pass, table, source_name or (table + ".csv"), "csv", data)
    except Exception:
        pass
    return {"table": table, "rows": n, "columns": cols, "ms": int((time.monotonic() - t0) * 1000)}


def import_xlsx(db, mysql_user, mysql_pass, table, data: bytes, has_header: bool, overwrite: bool,
                source_name: str = "") -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError_("服务器未安装 openpyxl")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ImportError_(f"Excel 文件解析失败：{e}")
    t0 = time.monotonic()
    try:
        ws = wb.worksheets[0]
        grid = []
        for row in ws.iter_rows(values_only=True):
            grid.append(list(row))
            if len(grid) > C.IMPORT_ROWS_MAX + 1:
                break
    finally:
        wb.close()
    cols, types, rows = _rows_from_grid(grid, has_header)
    n = _create_and_fill(db, mysql_user, mysql_pass, table, cols, types, rows, overwrite)
    try:
        store_source(db, mysql_user, mysql_pass, table, source_name or (table + ".xlsx"), "xlsx", data)
    except Exception:
        pass
    return {"table": table, "rows": n, "columns": cols, "ms": int((time.monotonic() - t0) * 1000)}


def import_sql(db, mysql_user, mysql_pass, data: bytes) -> dict:
    text = _decode(data)
    if len(text) > 2 * 1024 * 1024:
        raise ImportError_("SQL 脚本超过 2MB")
    return run_sql(db, mysql_user, mysql_pass, text)
